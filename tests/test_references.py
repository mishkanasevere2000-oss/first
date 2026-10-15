import sqlite3

from openpyxl import Workbook
import pytest

from program.database.connection import connect
from program.database.migrations import initialize, validate_database
from program.database.operations import Store
from program.database.references import HEADERS, create_template, import_reference


@pytest.fixture
def connection(tmp_path):
    path = tmp_path / 'finance.db'
    initialize(path)
    with connect(path) as connection:
        yield connection


def workbook(tmp_path, entity, rows):
    path = tmp_path / f'{entity}.xlsx'
    book = Workbook()
    book.active.append(HEADERS[entity])
    for row in rows:
        book.active.append(row)
    book.save(path)
    book.close()
    return path


def categories(tmp_path, rows=None):
    return workbook(tmp_path, 'categories', rows if rows is not None else [
        ('child', 'Подкатегория', 'root', 'ACTIVE'),
        ('root', 'Пример', None, 'ACTIVE'),
    ])


def test_category_import_repeat_and_audit(connection, tmp_path):
    path = categories(tmp_path)
    assert import_reference(connection, path, 'categories') == 2
    assert import_reference(connection, path, 'categories') == 0
    assert connection.execute('SELECT count(*) FROM categories').fetchone()[0] == 2
    assert connection.execute('SELECT count(*) FROM audit_log').fetchone()[0] == 2
    assert validate_database(connection) == 1


@pytest.mark.parametrize('rows', [
    [('a', 'Name', None, 'ACTIVE'), ('a', 'Duplicate ID', None, 'ACTIVE')],
    [('a', 'Имя', None, 'ACTIVE'), ('b', 'ИМЯ', None, 'ACTIVE')],
    [('a', 'Name', 'absent', 'ACTIVE')],
    [('a', 'Name', 'a', 'ACTIVE')],
    [('a', 'One', 'b', 'ACTIVE'), ('b', 'Two', 'a', 'ACTIVE')],
    [('a', None, None, 'ACTIVE')],
    [('a', 'Name', None, 'TYPO')],
    [(123, 'Numeric identifier', None, 'ACTIVE')],
    [('a', '=1+1', None, 'ACTIVE')],
])
def test_invalid_category_file_changes_nothing(connection, tmp_path, rows):
    assert import_reference(connection, categories(tmp_path), 'categories') == 2
    before = [tuple(row) for row in connection.execute('SELECT * FROM categories ORDER BY id')]
    with pytest.raises((ValueError, sqlite3.IntegrityError)):
        import_reference(connection, categories(tmp_path, rows), 'categories')
    assert [tuple(row) for row in connection.execute('SELECT * FROM categories ORDER BY id')] == before
    assert connection.execute('SELECT count(*) FROM audit_log').fetchone()[0] == 2


def test_category_update_preserves_references_and_missing_records(connection, tmp_path):
    import_reference(connection, categories(tmp_path), 'categories')
    store = Store(connection)
    store.create_import('i', 'fake.xlsx', 'Example')
    store.add_operation(id='o', import_id='i', bank='Example', account_id='a', operation_date='2026-01-01',
                        description='Synthetic', amount_kopecks=10, direction='OUT')
    store.record_manual_decision('o', operation_type='expense', include_in_report=True,
                                 category_id='root', subcategory_id='child', comment='Test decision')
    update = categories(tmp_path, [('root', 'Переименовано', None, 'ACTIVE')])
    assert import_reference(connection, update, 'categories') == 1
    assert store.get('categories', 'child') is not None
    assert store.get('classifications', 'o')['category_id'] == 'root'
    bad = categories(tmp_path, [('child', 'Подкатегория', None, 'ACTIVE')])
    with pytest.raises(ValueError):
        import_reference(connection, bad, 'categories')
    assert store.get('categories', 'child')['parent_id'] == 'root'


def test_category_cycle_with_existing_data(connection, tmp_path):
    import_reference(connection, categories(tmp_path), 'categories')
    with pytest.raises(ValueError):
        import_reference(connection, categories(tmp_path, [('root', 'Пример', 'child', 'ACTIVE')]), 'categories')


def test_casefold_conflict_with_retained_category(connection, tmp_path):
    import_reference(connection, categories(tmp_path), 'categories')
    with pytest.raises(ValueError):
        import_reference(connection, categories(tmp_path, [('other', 'ПРИМЕР', None, 'ACTIVE')]), 'categories')


def test_category_name_swap(connection, tmp_path):
    import_reference(connection, categories(tmp_path, [('a', 'One', None, 'ACTIVE'), ('b', 'Two', None, 'ACTIVE')]), 'categories')
    assert import_reference(connection, categories(tmp_path, [('a', 'Two', None, 'ACTIVE'), ('b', 'One', None, 'ACTIVE')]), 'categories') == 2


def test_accounts_repeat_update_archive_and_audit(connection, tmp_path):
    path = workbook(tmp_path, 'own_accounts', [('acc', 'Example Bank', 'DEMO-001', 'Example Owner', None, 'ACTIVE')])
    assert import_reference(connection, path, 'own_accounts') == 1
    assert import_reference(connection, path, 'own_accounts') == 0
    update = workbook(tmp_path, 'own_accounts', [('acc', 'Example Bank', 'DEMO-001', 'Example Owner', None, 'ARCHIVED')])
    assert import_reference(connection, update, 'own_accounts') == 1
    assert Store(connection).get('own_accounts', 'acc')['status'] == 'ARCHIVED'
    assert connection.execute('SELECT count(*) FROM audit_log').fetchone()[0] == 2


@pytest.mark.parametrize('rows', [
    [('a', 'Bank', 'DEMO', None, None, 'ACTIVE'), ('b', 'Bank', 'DEMO', None, None, 'ACTIVE')],
    [('a', 'Bank', 123, None, None, 'ACTIVE')],
    [('a', '', 'DEMO', None, None, 'ACTIVE')],
    [('a', 'Bank', 'DEMO', None, None, 'TYPO')],
])
def test_accounts_invalid(connection, tmp_path, rows):
    with pytest.raises(ValueError):
        import_reference(connection, workbook(tmp_path, 'own_accounts', rows), 'own_accounts')
    assert connection.execute('SELECT count(*) FROM own_accounts').fetchone()[0] == 0
    assert connection.execute('SELECT count(*) FROM audit_log').fetchone()[0] == 0


def test_error_mid_write_rolls_back_reference_and_audit(connection, tmp_path):
    original = Store._audit
    def fail(self, event, entity, identifier, before, after):
        if identifier == 'bad':
            raise RuntimeError('Injected failure after first update')
        original(self, event, entity, identifier, before, after)
    path = categories(tmp_path, [('a', 'One', None, 'ACTIVE'), ('bad', 'Two', None, 'ACTIVE')])
    from unittest.mock import patch
    with patch.object(Store, '_audit', fail):
        with pytest.raises(RuntimeError):
            import_reference(connection, path, 'categories')
    assert connection.execute('SELECT count(*) FROM categories').fetchone()[0] == 0
    assert connection.execute('SELECT count(*) FROM audit_log').fetchone()[0] == 0


def test_header_only_templates_and_no_overwrite(connection, tmp_path):
    for entity in HEADERS:
        path = tmp_path / f'{entity}.xlsx'
        create_template(path, entity)
        assert import_reference(connection, path, entity) == 0
        before = path.read_bytes()
        with pytest.raises(FileExistsError):
            create_template(path, entity)
        assert path.read_bytes() == before


def test_bad_headers_and_multiple_sheets(connection, tmp_path):
    path = tmp_path / 'bad.xlsx'
    book = Workbook()
    book.active.append(['wrong'])
    book.save(path)
    with pytest.raises(ValueError):
        import_reference(connection, path, 'categories')
    book.active.delete_rows(1)
    book.active.append(HEADERS['categories'])
    book.create_sheet('Unexpected')
    book.save(path)
    book.close()
    with pytest.raises(ValueError):
        import_reference(connection, path, 'categories')


def test_accounts_swap_is_atomic(connection, tmp_path):
    first = [('a', 'Bank', 'A', None, None, 'ACTIVE'), ('b', 'Bank', 'B', None, None, 'ACTIVE')]
    import_reference(connection, workbook(tmp_path, 'own_accounts', first), 'own_accounts')
    swapped = [('a', 'Bank', 'B', None, None, 'ACTIVE'), ('b', 'Bank', 'A', None, None, 'ACTIVE')]
    assert import_reference(connection, workbook(tmp_path, 'own_accounts', swapped), 'own_accounts') == 2
    assert Store(connection).get('own_accounts', 'a')['account'] == 'B'


def test_reference_update_has_verified_backup(connection, tmp_path):
    import_reference(connection, categories(tmp_path), 'categories')
    updated = categories(tmp_path, [('root', 'Новое имя', None, 'ACTIVE')])
    assert import_reference(connection, updated, 'categories', backup_dir=tmp_path / 'copies') == 1
    files = list((tmp_path / 'copies').glob('*.db'))
    assert len(files) == 1
    with connect(files[0], readonly=True) as backup:
        assert validate_database(backup) == 1
        assert Store(backup).get('categories', 'root')['name'] == 'Пример'
    assert import_reference(connection, updated, 'categories', backup_dir=tmp_path / 'copies') == 0
    assert len(list((tmp_path / 'copies').glob('*.db'))) == 1


def test_reference_backup_failure_prevents_updates(connection, tmp_path, monkeypatch):
    import program.database.references as module
    import_reference(connection, categories(tmp_path), 'categories')
    before = [tuple(row) for row in connection.execute('SELECT * FROM categories ORDER BY id')]
    def fail(*args):
        raise OSError('Backup failed')
    monkeypatch.setattr(module, 'verified_backup', fail)
    with pytest.raises(OSError):
        import_reference(connection, categories(tmp_path, [('root', 'Changed', None, 'ACTIVE')]), 'categories')
    assert [tuple(row) for row in connection.execute('SELECT * FROM categories ORDER BY id')] == before
    assert connection.execute('SELECT count(*) FROM audit_log').fetchone()[0] == 2
