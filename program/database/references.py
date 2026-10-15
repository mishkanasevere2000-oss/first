"""Validated XLSX reference import. No bank statement import is implemented."""
from pathlib import Path
import unicodedata
import uuid

from openpyxl import Workbook, load_workbook

from .connection import DEFAULT_DATABASE, PROJECT_ROOT, transaction
from .migrations import validate_database, verified_backup
from .operations import Store, require_text

HEADERS = {
    'categories': ('id', 'name', 'parent_id', 'status'),
    'own_accounts': ('id', 'bank', 'account', 'owner_name', 'phone', 'status'),
}


def name_key(value):
    return unicodedata.normalize('NFKC', value.strip()).casefold()


def read_reference(path, entity):
    if entity not in HEADERS:
        raise ValueError('Поддерживаются только categories и own_accounts.')
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError('Справочник должен содержать ровно один лист.')
        rows = workbook.active.iter_rows()
        header = next(rows, ())
        if tuple(cell.value for cell in header) != HEADERS[entity]:
            raise ValueError('Неверные колонки справочника: ' + ', '.join(HEADERS[entity]))
        result = {}
        for number, cells in enumerate(rows, start=2):
            if all(cell.value is None for cell in cells):
                continue
            if any(cell.data_type in {'f', 'e'} for cell in cells):
                raise ValueError(f'Строка {number}: формулы и ошибки Excel не допускаются.')
            values = [cell.value for cell in cells]
            if len(values) != len(HEADERS[entity]):
                raise ValueError(f'Строка {number}: неверное число колонок.')
            record = dict(zip(HEADERS[entity], values))
            optional = {'parent_id', 'owner_name', 'phone'}
            for key, value in record.items():
                if value is None and key in optional:
                    continue
                record[key] = require_text(value, f'Строка {number}, {key}')
            if record['status'] not in {'ACTIVE', 'ARCHIVED'}:
                raise ValueError(f'Строка {number}: статус должен быть ACTIVE или ARCHIVED.')
            if record['id'] in result:
                raise ValueError(f'Строка {number}: повторный id.')
            result[record['id']] = record
        return result
    finally:
        workbook.close()


def validate_merged(entity, records):
    if entity == 'categories':
        seen = set()
        for identifier, row in records.items():
            key = (row['parent_id'], name_key(row['name']))
            if key in seen:
                raise ValueError('Повторное имя категории внутри одного родителя.')
            seen.add(key)
            visited = set()
            node = identifier
            while node is not None:
                if node in visited:
                    raise ValueError('Обнаружен цикл категорий.')
                if node not in records:
                    raise ValueError('Родительская категория не найдена.')
                visited.add(node)
                node = records[node]['parent_id']
    else:
        keys = [(row['bank'], row['account']) for row in records.values()]
        if len(keys) != len(set(keys)):
            raise ValueError('Счёт уже существует под другим id.')


def normalized_record(entity, record):
    values = dict(record)
    if entity == 'categories':
        values['name_key'] = name_key(values['name'])
    return values


def validate_current_merge(connection, entity, incoming):
    existing = {row['id']: dict(row) for row in connection.execute(f'SELECT * FROM {entity}')}
    merged = {**existing, **incoming}
    validate_merged(entity, merged)
    if entity == 'categories':
        for table in ('classifications', 'manual_overrides'):
            for row in connection.execute(f'SELECT category_id,subcategory_id FROM {table} WHERE subcategory_id IS NOT NULL'):
                if merged[row[1]]['parent_id'] != row[0]:
                    raise ValueError('Перенос категории нарушает существующее или историческое решение.')
    return existing


def import_reference(connection, path, entity, *, backup_dir=None):
    incoming = read_reference(path, entity)
    store = Store(connection)
    # Back up before changing any existing reference. Repeats produce no backup.
    before = validate_current_merge(connection, entity, incoming)
    changed_existing = any(identifier in before and before[identifier] != normalized_record(entity, row)
                           for identifier, row in incoming.items())
    if changed_existing:
        if backup_dir is None:
            database_path = Path(connection.execute('PRAGMA database_list').fetchone()[2])
            backup_dir = (PROJECT_ROOT if database_path == DEFAULT_DATABASE else database_path.parent) / 'backup'
        verified_backup(connection, backup_dir, validate_database)
    changed = 0
    with transaction(connection):
        existing = validate_current_merge(connection, entity, incoming)
        # Only modify changed rows. This preserves the no-write repeat guarantee.
        updates = {identifier: normalized_record(entity, row) for identifier, row in incoming.items()
                   if existing.get(identifier) != normalized_record(entity, row)}
        for identifier in updates:
            if identifier in existing:
                if entity == 'categories':
                    connection.execute('UPDATE categories SET name_key=? WHERE id=?',
                                       ('reserved:' + uuid.uuid4().hex, identifier))
                else:
                    connection.execute('UPDATE own_accounts SET bank=? WHERE id=?',
                                       ('reserved:' + uuid.uuid4().hex, identifier))
        if entity == 'categories':
            connection.execute('PRAGMA defer_foreign_keys=ON')
        for identifier, values in updates.items():
            columns = ','.join(values)
            assignments = ','.join(f'{key}=excluded.{key}' for key in values if key != 'id')
            connection.execute(
                f"INSERT INTO {entity}({columns}) VALUES ({','.join('?' for _ in values)}) "
                f'ON CONFLICT(id) DO UPDATE SET {assignments}', tuple(values.values()))
            store._audit('REFERENCE_IMPORT', entity, identifier, existing.get(identifier), values)
            changed += 1
    return changed


def create_template(path, entity):
    if entity not in HEADERS:
        raise ValueError('Неизвестный справочник.')
    path = Path(path)
    workbook = Workbook()
    workbook.active.title = entity
    workbook.active.append(HEADERS[entity])
    workbook.active.freeze_panes = 'A2'
    # Never overwrite a user file.
    try:
        with path.open('xb') as output:
            workbook.save(output)
    finally:
        workbook.close()
